"""Import des offres des soumissionnaires et analyse comparative multi-lots.

Un fichier d'offre peut couvrir un nombre quelconque de lots : chaque ligne est
rattachée à son poste par le code de poste, quel que soit l'onglet où elle se
trouve. Un lot est considéré comme remis dès qu'un de ses postes est chiffré.
"""

import io
import re
import statistics
from pathlib import Path

from openpyxl import load_workbook

from .db import lots_du_projet

ENTETES = {
    "code": ("code", "n°", "no", "poste", "article"),
    "quantite": ("quantité", "quantite", "qté", "qte", "qt"),
    "pu": ("prix unitaire", "p.u.", "pu", "prix unit."),
    "total": ("total", "montant", "prix total"),
}
LIBELLE_SOUMISSIONNAIRE = ("soumissionnaire", "entreprise", "entrepreneur")
TOLERANCE_CALCUL = 0.01


class OffreIllisible(ValueError):
    pass


def normaliser_code(valeur) -> str:
    texte = str(valeur).strip().lower().rstrip(".")
    return re.sub(r"\s+", "", texte)


def lire_nombre(valeur):
    if valeur is None or valeur == "":
        return None
    if isinstance(valeur, (int, float)):
        return float(valeur)
    texte = re.sub(r"[^\d,.\-]", "", str(valeur))
    if not texte:
        return None
    if "," in texte:
        texte = texte.replace(".", "").replace(",", ".")
    try:
        return float(texte)
    except ValueError:
        return None


def _colonnes(ligne) -> dict:
    trouvees = {}
    for idx, cellule in enumerate(ligne):
        if not isinstance(cellule, str):
            continue
        texte = cellule.strip().lower()
        for cle, variantes in ENTETES.items():
            if cle not in trouvees and any(
                texte == v or (len(v) > 3 and texte.startswith(v)) for v in variantes
            ):
                trouvees[cle] = idx
                break
    return trouvees if "code" in trouvees and ("pu" in trouvees or "total" in trouvees) else {}


def _entreprise(classeur) -> str | None:
    for feuille in classeur.worksheets:
        for ligne in feuille.iter_rows(max_row=20, values_only=True):
            for idx, cellule in enumerate(ligne):
                if isinstance(cellule, str) and cellule.strip().lower().rstrip(" :") in LIBELLE_SOUMISSIONNAIRE:
                    suite = [c for c in ligne[idx + 1:] if c not in (None, "")]
                    if suite:
                        return str(suite[0]).strip()
    return None


def lire_offre(contenu: bytes, nom_fichier: str, postes: list[dict]) -> dict:
    """Lit un classeur d'offre et rattache chaque prix au poste correspondant."""
    try:
        classeur = load_workbook(io.BytesIO(contenu), data_only=True, read_only=True)
    except Exception as exc:
        raise OffreIllisible(f"{nom_fichier} : fichier Excel illisible") from exc

    par_code = {normaliser_code(p["code"]): p for p in postes}
    prix, inconnus = {}, []
    for feuille in classeur.worksheets:
        colonnes = {}
        for ligne in feuille.iter_rows(values_only=True):
            if not colonnes:
                colonnes = _colonnes(ligne)
                continue
            code = ligne[colonnes["code"]] if colonnes["code"] < len(ligne) else None
            if code in (None, ""):
                continue
            valeurs = {cle: lire_nombre(ligne[i]) if i < len(ligne) else None
                       for cle, i in colonnes.items() if cle != "code"}
            if valeurs.get("pu") is None and valeurs.get("total") is None:
                continue
            poste = par_code.get(normaliser_code(code))
            if poste is None:
                inconnus.append(str(code))
                continue
            prix[poste["id"]] = {
                "quantite": valeurs.get("quantite"),
                "pu": valeurs.get("pu"),
                "total": valeurs.get("total"),
            }
    entreprise = _entreprise(classeur)
    classeur.close()
    if not prix:
        raise OffreIllisible(
            f"{nom_fichier} : aucun prix reconnu (colonnes « Code » et « Prix unitaire » attendues)")
    return {
        "entreprise": entreprise or Path(nom_fichier).stem,
        "prix": prix,
        "inconnus": inconnus,
    }


def enregistrer_offre(conn, projet_id: int, offre: dict, nom_fichier: str) -> int:
    """Enregistre l'offre ; une nouvelle remise de la même entreprise remplace l'ancienne."""
    conn.execute("DELETE FROM offre WHERE projet_id = ? AND entreprise = ?",
                 (projet_id, offre["entreprise"]))
    offre_id = conn.execute(
        "INSERT INTO offre (projet_id, entreprise, fichier) VALUES (?, ?, ?)",
        (projet_id, offre["entreprise"], nom_fichier),
    ).lastrowid
    conn.executemany(
        "INSERT INTO offre_prix (offre_id, poste_id, quantite, pu, total) VALUES (?, ?, ?, ?, ?)",
        [(offre_id, pid, v["quantite"], v["pu"], v["total"]) for pid, v in offre["prix"].items()],
    )
    return offre_id


def _pu_effectif(poste, ligne):
    if ligne["pu"] is not None:
        return ligne["pu"]
    if ligne["total"] is None:
        return None
    if poste["mode"] == "FF":
        return ligne["total"]
    return ligne["total"] / poste["quantite"] if poste["quantite"] else None


def _montant(poste, pu):
    if pu is None or poste["mode"] == "PM":
        return 0.0
    return pu if poste["mode"] == "FF" else pu * poste["quantite"]


def analyser(conn, projet_id: int, seuil_anormal: float = 0.25) -> dict:
    """Tableau comparatif par lot, contrôles et synthèse d'attribution.

    Les montants sont recalculés sur les quantités du métré (montant corrigé) ;
    le montant déclaré par le soumissionnaire est conservé pour comparaison.
    """
    offres = [dict(o) for o in conn.execute(
        "SELECT * FROM offre WHERE projet_id = ? ORDER BY entreprise", (projet_id,))]
    saisies = {}
    for r in conn.execute(
        "SELECT op.* FROM offre_prix op JOIN offre o ON o.id = op.offre_id WHERE o.projet_id = ?",
        (projet_id,),
    ):
        saisies[(r["offre_id"], r["poste_id"])] = dict(r)

    lots = []
    for lot in lots_du_projet(conn, projet_id):
        remises = [o for o in offres
                   if any((o["id"], p["id"]) in saisies for p in lot["postes"])]
        lignes = []
        totaux = {o["id"]: {"corrige": 0.0, "declare": 0.0, "alertes": 0} for o in remises}
        for poste in lot["postes"]:
            cellules = {}
            for o in remises:
                saisie = saisies.get((o["id"], poste["id"]))
                pu = _pu_effectif(poste, saisie) if saisie else None
                corrige = _montant(poste, pu)
                alertes = []
                if pu is None and poste["mode"] != "PM":
                    alertes.append("prix manquant")
                if saisie and saisie["total"] is not None and saisie["pu"] is not None:
                    if abs(saisie["total"] - corrige) > TOLERANCE_CALCUL + 0.001 * abs(corrige):
                        alertes.append("erreur de calcul")
                if (saisie and saisie["quantite"] is not None and poste["mode"] in ("QF", "QP")
                        and abs(saisie["quantite"] - poste["quantite"]) > 1e-6):
                    alertes.append("quantité modifiée")
                cellules[o["id"]] = {"pu": pu, "montant": corrige, "alertes": alertes}
                totaux[o["id"]]["corrige"] += corrige
                totaux[o["id"]]["declare"] += (saisie["total"] if saisie and saisie["total"] is not None
                                               else corrige)
            _marquer_prix_anormaux(poste, cellules, seuil_anormal)
            prix = [c["montant"] for c in cellules.values() if c["pu"] is not None]
            for oid, c in cellules.items():
                c["moins_disant"] = bool(prix) and c["pu"] is not None and c["montant"] == min(prix)
                totaux[oid]["alertes"] += len(c["alertes"])
            lignes.append({"poste": poste, "estimation": _montant(poste, poste["prix_unitaire"]),
                           "offres": cellules})

        classement = sorted(
            ({"offre": o, **totaux[o["id"]],
              "ecart": _ecart(totaux[o["id"]]["corrige"], lot["total"])} for o in remises),
            key=lambda t: t["corrige"],
        )
        lots.append({"lot": lot, "offres": remises, "lignes": lignes,
                     "totaux": totaux, "classement": classement})

    return {"offres": offres, "lots": lots, "synthese": _synthese(offres, lots)}


def _marquer_prix_anormaux(poste, cellules, seuil):
    pus = [c["pu"] for c in cellules.values() if c["pu"] is not None]
    if len(pus) >= 3:
        reference = statistics.median(pus)
    elif poste["prix_unitaire"]:
        reference = poste["prix_unitaire"]
    else:
        return
    if not reference:
        return
    for c in cellules.values():
        if c["pu"] is None:
            continue
        ecart = (c["pu"] - reference) / reference
        if ecart > seuil:
            c["alertes"].append(f"prix anormalement élevé ({ecart:+.0%})")
        elif ecart < -seuil:
            c["alertes"].append(f"prix anormalement bas ({ecart:+.0%})")


def _ecart(valeur, reference):
    return (valeur - reference) / reference if reference else None


def _synthese(offres, lots):
    par_lot = []
    for lot in lots:
        meilleure = lot["classement"][0] if lot["classement"] else None
        par_lot.append({"lot": lot["lot"], "meilleure": meilleure})
    attribution = sum(l["meilleure"]["corrige"] for l in par_lot if l["meilleure"])
    lots_ids = {l["lot"]["id"] for l in lots}
    globales = []
    for o in offres:
        couverts = {l["lot"]["id"] for l in lots if o["id"] in l["totaux"]}
        globales.append({
            "offre": o,
            "lots": len(couverts),
            "complete": couverts == lots_ids,
            "total": sum(l["totaux"][o["id"]]["corrige"] for l in lots if o["id"] in l["totaux"]),
        })
    return {"par_lot": par_lot, "attribution_par_lot": attribution,
            "lots_sans_offre": [l["lot"] for l in par_lot if not l["meilleure"]],
            "globales": sorted(globales, key=lambda g: (not g["complete"], g["total"]))}
