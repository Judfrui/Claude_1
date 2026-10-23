import io
import re

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import MODES, lots_du_projet

FORMAT_MONTANT = '#,##0.00'
GRAS = Font(bold=True)
A_REMPLIR = PatternFill("solid", fgColor="FFF2CC")
MOINS_DISANT = PatternFill("solid", fgColor="D9EAD3")
ALERTE = Font(color="C00000")


def nombre_fr(valeur: float, decimales: int = 2) -> str:
    texte = f"{valeur:,.{decimales}f}"
    return texte.replace(",", " ").replace(".", ",")


def _projet(conn, projet_id):
    return conn.execute("SELECT * FROM projet WHERE id = ?", (projet_id,)).fetchone()


def _nom_onglet(lot) -> str:
    return re.sub(r"[\[\]:*?/\\]", "-", f"Lot {lot['code']}")[:31]


def csc_docx(conn, projet_id: int) -> bytes:
    """Cahier spécial des charges — clauses techniques, par lot et par poste."""
    projet = _projet(conn, projet_id)
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Cahier spécial des charges", level=0)
    doc.add_paragraph("Clauses techniques")
    infos = doc.add_table(rows=0, cols=2)
    for libelle, valeur in (("Projet", projet["nom"]), ("Référence", projet["reference"]),
                            ("Maître de l'ouvrage", projet["maitre_ouvrage"]),
                            ("Adresse du chantier", projet["adresse"])):
        if valeur:
            cellules = infos.add_row().cells
            cellules[0].text, cellules[1].text = libelle, valeur

    for lot in lots_du_projet(conn, projet_id):
        doc.add_heading(f"Lot {lot['code']} — {lot['titre']}", level=1)
        for poste in lot["postes"]:
            doc.add_heading(f"{poste['code']}  {poste['titre']}", level=2)
            for paragraphe in filter(None, (poste["description"] or "").splitlines()):
                doc.add_paragraph(paragraphe)
            mesure = f"Unité : {poste['unite'] or '—'} — Mode de métré : {poste['mode']} ({MODES[poste['mode']]})"
            if poste["mode"] in ("QF", "QP"):
                mesure += f" — Quantité : {nombre_fr(poste['quantite'])}"
            doc.add_paragraph().add_run(mesure).italic = True

    sortie = io.BytesIO()
    doc.save(sortie)
    return sortie.getvalue()


def metre_xlsx(conn, projet_id: int, avec_prix: bool) -> bytes:
    """Métré récapitulatif : vierge (à remettre aux soumissionnaires) ou estimatif.

    Un seul classeur couvre tous les lots, si bien qu'un soumissionnaire peut
    remettre prix pour plusieurs lots dans un seul fichier.
    """
    projet = _projet(conn, projet_id)
    lots = lots_du_projet(conn, projet_id)
    classeur = Workbook()
    garde = classeur.active
    garde.title = "Soumission"
    garde.append(["Projet", projet["nom"]])
    garde.append(["Référence", projet["reference"]])
    garde.append(["Maître de l'ouvrage", projet["maitre_ouvrage"]])
    garde.append([])
    garde.append(["Soumissionnaire :", ""])
    garde["B5"].fill = A_REMPLIR
    garde.append([])
    garde.append(["Complétez les prix unitaires (cellules jaunes) des lots pour lesquels vous remettez offre. "
                  "Laissez vides les lots non remis. Ne modifiez pas la colonne « Code »."])
    for cellule in ("A1", "A2", "A3", "A5"):
        garde[cellule].font = GRAS
    garde.column_dimensions["A"].width = 22
    garde.column_dimensions["B"].width = 50

    totaux = []
    for lot in lots:
        feuille = classeur.create_sheet(_nom_onglet(lot))
        feuille.append([f"Lot {lot['code']} — {lot['titre']}"])
        feuille["A1"].font = Font(bold=True, size=12)
        feuille.append(["Code", "Désignation", "Unité", "Mode", "Quantité", "Prix unitaire", "Total"])
        for cellule in feuille[2]:
            cellule.font = GRAS
        for poste in lot["postes"]:
            r = feuille.max_row + 1
            quantite = 1 if poste["mode"] == "FF" else poste["quantite"]
            pu = poste["prix_unitaire"] if avec_prix else None
            total = "PM" if poste["mode"] == "PM" else f"=E{r}*F{r}"
            feuille.append([poste["code"], poste["titre"], poste["unite"], poste["mode"],
                            quantite, pu, total])
            feuille[f"E{r}"].number_format = FORMAT_MONTANT
            feuille[f"F{r}"].number_format = FORMAT_MONTANT
            feuille[f"G{r}"].number_format = FORMAT_MONTANT
            if poste["mode"] != "PM":
                feuille[f"F{r}"].fill = A_REMPLIR
        r = feuille.max_row + 1
        feuille.append([None, f"Total lot {lot['code']}", None, None, None, None,
                        f"=SUM(G3:G{r - 1})" if r > 3 else 0])
        feuille[f"B{r}"].font = GRAS
        feuille[f"G{r}"].font = GRAS
        feuille[f"G{r}"].number_format = FORMAT_MONTANT
        totaux.append((lot, f"='{feuille.title}'!G{r}"))
        for colonne, largeur in zip("ABCDEFG", (10, 60, 8, 7, 12, 14, 16)):
            feuille.column_dimensions[colonne].width = largeur
        feuille["B3"].alignment = Alignment(wrap_text=True)

    recap = classeur.create_sheet("Récapitulatif", 1)
    recap.append(["Lot", "Intitulé", "Total HTVA"])
    for cellule in recap[1]:
        cellule.font = GRAS
    for lot, formule in totaux:
        recap.append([lot["code"], lot["titre"], formule])
    n = recap.max_row
    recap.append([None, "Total HTVA", f"=SUM(C2:C{n})"])
    recap.append([None, "Taux de TVA", 0.21])
    recap.append([None, "Total TVAC", f"=C{n + 1}*(1+C{n + 2})"])
    for r in range(2, recap.max_row + 1):
        recap[f"C{r}"].number_format = FORMAT_MONTANT
    recap[f"C{n + 2}"].number_format = "0%"
    recap.column_dimensions["B"].width = 50
    recap.column_dimensions["C"].width = 18

    sortie = io.BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()


def analyse_xlsx(conn, projet_id: int, analyse: dict) -> bytes:
    """Rapport d'analyse des offres : synthèse et comparatif détaillé par lot."""
    projet = _projet(conn, projet_id)
    classeur = Workbook()
    synthese = classeur.active
    synthese.title = "Synthèse"
    synthese.append([f"Analyse des offres — {projet['nom']}"])
    synthese["A1"].font = Font(bold=True, size=12)
    synthese.append([])
    synthese.append(["Lot", "Intitulé", "Estimation", "Moins-disant", "Montant corrigé", "Écart / estimation"])
    for cellule in synthese[3]:
        cellule.font = GRAS
    for entree in analyse["synthese"]["par_lot"]:
        meilleure = entree["meilleure"]
        synthese.append([
            entree["lot"]["code"], entree["lot"]["titre"], entree["lot"]["total"],
            meilleure["offre"]["entreprise"] if meilleure else "Aucune offre",
            meilleure["corrige"] if meilleure else None,
            meilleure["ecart"] if meilleure else None,
        ])
    synthese.append([None, "Attribution lot par lot", None, None,
                     analyse["synthese"]["attribution_par_lot"]])
    synthese.append([])
    synthese.append(["Soumissionnaire", "Lots remis", "Offre complète", "Total corrigé"])
    for cellule in synthese[synthese.max_row]:
        cellule.font = GRAS
    for g in analyse["synthese"]["globales"]:
        synthese.append([g["offre"]["entreprise"], g["lots"], "oui" if g["complete"] else "non", g["total"]])
    for ligne in synthese.iter_rows(min_row=4):
        for cellule in ligne:
            if isinstance(cellule.value, float):
                cellule.number_format = FORMAT_MONTANT
        if isinstance(ligne[5].value, float):
            ligne[5].number_format = "+0.0%;-0.0%"
    for colonne, largeur in zip("ABCDEF", (24, 45, 16, 28, 16, 16)):
        synthese.column_dimensions[colonne].width = largeur

    for lot in analyse["lots"]:
        feuille = classeur.create_sheet(_nom_onglet(lot["lot"]))
        entete = ["Code", "Désignation", "Unité", "Quantité", "PU estimé", "Estimation"]
        for o in lot["offres"]:
            entete += [f"{o['entreprise']} — PU", "Montant", "Alertes"]
        feuille.append(entete)
        for cellule in feuille[1]:
            cellule.font = GRAS
            cellule.alignment = Alignment(wrap_text=True)
        for ligne in lot["lignes"]:
            p = ligne["poste"]
            valeurs = [p["code"], p["titre"], p["unite"], p["quantite"], p["prix_unitaire"], ligne["estimation"]]
            for o in lot["offres"]:
                c = ligne["offres"][o["id"]]
                valeurs += [c["pu"], c["montant"], ", ".join(c["alertes"])]
            feuille.append(valeurs)
            r = feuille.max_row
            for i, o in enumerate(lot["offres"]):
                col = 7 + 3 * i
                c = ligne["offres"][o["id"]]
                if c["moins_disant"]:
                    feuille.cell(r, col + 1).fill = MOINS_DISANT
                if c["alertes"]:
                    feuille.cell(r, col + 2).font = ALERTE
        feuille.append([None, "Total corrigé", None, None, None, lot["lot"]["total"]]
                       + sum(([None, lot["totaux"][o["id"]]["corrige"], None] for o in lot["offres"]), []))
        feuille.append([None, "Total déclaré", None, None, None, None]
                       + sum(([None, lot["totaux"][o["id"]]["declare"], None] for o in lot["offres"]), []))
        for ligne in feuille.iter_rows(min_row=2):
            for cellule in ligne:
                if isinstance(cellule.value, float):
                    cellule.number_format = FORMAT_MONTANT
        feuille.column_dimensions["B"].width = 45
        for i in range(len(lot["offres"])):
            feuille.column_dimensions[get_column_letter(9 + 3 * i)].width = 30

    sortie = io.BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()
