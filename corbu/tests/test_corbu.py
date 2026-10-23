import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from corbu import db, offres
from corbu.app import creer_app
from corbu.formules import FormuleInvalide, evaluer


@pytest.fixture
def client(tmp_path):
    return TestClient(creer_app(str(tmp_path / "test.db")))


@pytest.fixture
def conn(client, tmp_path):
    c = db.connecter(str(tmp_path / "test.db"))
    yield c
    c.close()


def test_formules():
    assert evaluer("2 x 3,5") == 7
    assert evaluer("12,40 x 2,80 - 2 x 1,20 x 2,10") == pytest.approx(29.68)
    assert evaluer("(4 + 6) / 4") == 2.5
    assert evaluer("-1,5") == -1.5
    for invalide in ("__import__('os')", "2 x", "1 / 0", "a + 1", "9 ** 99"):
        with pytest.raises(FormuleInvalide):
            evaluer(invalide)


def test_lire_nombre():
    assert offres.lire_nombre("1.234,56 €") == 1234.56
    assert offres.lire_nombre("45,5") == 45.5
    assert offres.lire_nombre(12) == 12.0
    assert offres.lire_nombre("") is None
    assert offres.lire_nombre("néant") is None


def _projet(client, conn):
    r = client.post("/projets", data={"nom": "Maison Dupont", "reference": "2026-014"})
    projet_id = int(r.url.path.rsplit("/", 1)[1])
    ids = [conn.execute("SELECT id FROM bib_article WHERE code = ?", (c,)).fetchone()[0]
           for c in ("02.02", "02.03", "04.01", "00.01")]
    client.post(f"/projets/{projet_id}/articles", data={"articles": ids})
    postes = {p["code"]: p["id"] for p in db.postes_du_projet(conn, projet_id)}
    for code, formule, pu in (("02.02", "10 x 0,6 x 0,8", "40"), ("02.03", "10 x 0,6 x 0,3", "200"),
                              ("04.01", "2 x 12,40 x 2,80", "60"), ("00.01", "", "3000")):
        if formule:
            client.post(f"/postes/{postes[code]}/metre", data={"formule": formule})
        poste = conn.execute("SELECT * FROM poste WHERE id = ?", (postes[code],)).fetchone()
        client.post(f"/postes/{postes[code]}", data={"titre": poste["titre"], "mode": poste["mode"],
                                                      "unite": poste["unite"], "prix_unitaire": pu})
    return projet_id


def _remplir(metre: bytes, entreprise: str, prix: dict) -> bytes:
    classeur = load_workbook(io.BytesIO(metre))
    classeur["Soumission"]["B5"] = entreprise
    for feuille in classeur.worksheets:
        for ligne in feuille.iter_rows(min_row=3):
            if ligne[0].value in prix:
                ligne[5].value = prix[ligne[0].value]
    sortie = io.BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()


def _fichier_libre(lignes) -> bytes:
    """Offre remise dans un tableau personnel de l'entrepreneur (colonnes réordonnées)."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.append(["Offre de prix — Entreprise Gamma"])
    feuille.append(["Désignation", "N°", "Qté", "P.U.", "Prix total"])
    for ligne in lignes:
        feuille.append(ligne)
    sortie = io.BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()


def test_exports(client, conn):
    projet_id = _projet(client, conn)
    csc = Document(io.BytesIO(client.get(f"/projets/{projet_id}/csc.docx").content))
    textes = [p.text for p in csc.paragraphs]
    assert "Lot 02 — Terrassements et fondations" in textes
    assert any("Quantité : 69,44" in t for t in textes)

    classeur = load_workbook(io.BytesIO(client.get(f"/projets/{projet_id}/metre.xlsx").content))
    assert classeur.sheetnames == ["Soumission", "Récapitulatif", "Lot 00", "Lot 02", "Lot 04"]
    assert classeur["Lot 02"]["A3"].value == "02.02"
    assert classeur["Lot 02"]["F3"].value is None
    estimation = load_workbook(io.BytesIO(client.get(f"/projets/{projet_id}/metre.xlsx?prix=true").content))
    assert estimation["Lot 02"]["F3"].value == 40


def test_import_groupe_multi_lots(client, conn):
    projet_id = _projet(client, conn)
    metre = client.get(f"/projets/{projet_id}/metre.xlsx").content
    alpha = _remplir(metre, "Alpha SA", {"00.01": 2800, "02.02": 38, "02.03": 210, "04.01": 55})
    beta = _remplir(metre, "Bêta SRL", {"02.02": 45, "02.03": 220})
    gamma = _fichier_libre([["Maçonnerie", "04.01", 69.44, 70, 4000], ["Inconnu", "99.99", 1, 5, 5]])

    r = client.post(f"/projets/{projet_id}/offres", files=[
        ("fichiers", ("alpha.xlsx", alpha)), ("fichiers", ("beta.xlsx", beta)),
        ("fichiers", ("Gamma.xlsx", gamma)), ("fichiers", ("note.xlsx", b"pas un classeur")),
    ])
    assert r.status_code == 200
    assert "Codes inconnus ignorés : 99.99" in r.text
    assert "note.xlsx : fichier Excel illisible" in r.text

    analyse = offres.analyser(conn, projet_id)
    noms = {o["entreprise"] for o in analyse["offres"]}
    assert noms == {"Alpha SA", "Bêta SRL", "Gamma"}

    lots = {l["lot"]["code"]: l for l in analyse["lots"]}
    assert [c["offre"]["entreprise"] for c in lots["02"]["classement"]] == ["Alpha SA", "Bêta SRL"]
    assert lots["02"]["classement"][0]["corrige"] == pytest.approx(4.8 * 38 + 1.8 * 210)
    assert [c["offre"]["entreprise"] for c in lots["04"]["classement"]] == ["Alpha SA", "Gamma"]
    assert [c["offre"]["entreprise"] for c in lots["00"]["classement"]] == ["Alpha SA"]

    gamma_id = next(o["id"] for o in analyse["offres"] if o["entreprise"] == "Gamma")
    ligne = lots["04"]["lignes"][0]["offres"][gamma_id]
    assert "erreur de calcul" in ligne["alertes"]
    assert lots["04"]["totaux"][gamma_id]["declare"] == 4000
    assert lots["04"]["totaux"][gamma_id]["corrige"] == pytest.approx(69.44 * 70)

    globales = {g["offre"]["entreprise"]: g for g in analyse["synthese"]["globales"]}
    assert globales["Alpha SA"]["complete"] and not globales["Bêta SRL"]["complete"]
    assert analyse["synthese"]["globales"][0]["offre"]["entreprise"] == "Alpha SA"

    rapport = load_workbook(io.BytesIO(client.get(f"/projets/{projet_id}/analyse.xlsx").content))
    assert rapport.sheetnames == ["Synthèse", "Lot 00", "Lot 02", "Lot 04"]


def test_reimport_remplace_offre(client, conn):
    projet_id = _projet(client, conn)
    metre = client.get(f"/projets/{projet_id}/metre.xlsx").content
    client.post(f"/projets/{projet_id}/offres",
                files=[("fichiers", ("a.xlsx", _remplir(metre, "Alpha SA", {"02.02": 38})))])
    client.post(f"/projets/{projet_id}/offres",
                files=[("fichiers", ("a2.xlsx", _remplir(metre, "Alpha SA", {"04.01": 50})))])
    analyse = offres.analyser(conn, projet_id)
    assert len(analyse["offres"]) == 1
    assert [l["lot"]["code"] for l in analyse["lots"] if l["offres"]] == ["04"]


def test_prix_anormaux(client, conn):
    projet_id = _projet(client, conn)
    metre = client.get(f"/projets/{projet_id}/metre.xlsx").content
    fichiers = [("fichiers", (f"{n}.xlsx", _remplir(metre, n, {"02.02": pu})))
                for n, pu in (("A", 40), ("B", 42), ("C", 80), ("D", 15))]
    client.post(f"/projets/{projet_id}/offres", files=fichiers)
    analyse = offres.analyser(conn, projet_id, seuil_anormal=0.25)
    lot = next(l for l in analyse["lots"] if l["lot"]["code"] == "02")
    alertes = {o["entreprise"]: lot["lignes"][0]["offres"][o["id"]]["alertes"] for o in lot["offres"]}
    assert any("élevé" in a for a in alertes["C"])
    assert any("bas" in a for a in alertes["D"])
    assert not any("anormalement" in a for a in alertes["A"])


def test_pages(client, conn):
    projet_id = _projet(client, conn)
    for url in ("/", f"/projets/{projet_id}", f"/projets/{projet_id}/offres", "/bibliotheque"):
        assert client.get(url).status_code == 200
    assert client.get("/projets/999").status_code == 404
    r = client.post(f"/postes/1/metre", data={"formule": "import os"})
    assert "Formule invalide" in r.text


def test_bibliotheque_aller_retour(client, conn):
    exportee = client.get("/bibliotheque.json").json()
    assert exportee["lots"][0]["code"] == "00"
    r = client.post("/bibliotheque/import",
                    files={"fichier_json": ("b.json", b'{"lots":[{"code":"99","titre":"Divers","articles":'
                                                      b'[{"code":"99.01","titre":"Essai","mode":"XX"}]}]}')})
    assert "invalide" in r.text
    assert conn.execute("SELECT COUNT(*) FROM bib_lot WHERE code = '99'").fetchone()[0] == 0
